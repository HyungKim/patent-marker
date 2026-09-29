"""
tests/test_prehl.py ─ 원본 형광펜 정리 점검 (LLM 없이, 30초)
=====================================================================
작성자가 미리 칠해 둔 형광펜이 든 문서를 만들어
  PPTX : 지우고 시작 → 마킹본에는 도구 표시만 남고, 검토 반영에 가짜 누락이 생기지 않는지
         형광펜을 지우면 안 보이게 될 글자(진한 형광펜 + 흰 글자)를 보정하는지, 원본 파일이 그대로인지
         끄면(--keep-highlights) 그대로 남고 경고가 나오는지, 마킹본을 다시 분석해도 깨끗한지
  PDF  : 뷰어로 칠한 형광펜 주석은 지우고, 쪽에 '그림'으로 굳은 형광펜은 바탕색으로 덮는지
         표 칸·도형 같은 채움과 밝은 글자는 건드리지 않는지
를 확인합니다. 모델 자리는 tests/mock_llm.py 의 가짜 Ollama 가 대신합니다.

    실행:  python tests/test_prehl.py
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

for _s in (sys.stdout, sys.stderr):                    # Windows 콘솔(cp949) 대비
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:  # noqa: BLE001
        pass

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
sys.path.insert(0, str(ROOT / "tools"))

TMP = Path(tempfile.mkdtemp(prefix="pm-prehl-test-"))
os.environ["PM_REVIEW_DIR"] = str(TMP / "review_data")
os.environ["PM_INPUT_DIR"] = str(TMP / "input")
os.environ["PM_OUTPUT_DIR"] = str(TMP / "output")
for _k in ("PM_STRIP_HL", "PM_PDF_COVER"):             # 점검은 기본값(켜짐) 기준
    os.environ.pop(_k, None)

import mock_llm  # noqa: E402

MOCK_HOST = mock_llm.start()
os.environ["PM_OLLAMA_HOST"] = MOCK_HOST

from pptx import Presentation                                        # noqa: E402
from pptx.dml.color import RGBColor                                  # noqa: E402
from pptx.oxml.ns import qn                                          # noqa: E402
from pypdf import PdfReader, PdfWriter                               # noqa: E402
from pypdf.annotations import Highlight, Text                        # noqa: E402

import make_sample_pdf                                               # noqa: E402
from app import analyze, config, extract, mark, pdfdoc, pipeline, prehl, review  # noqa: E402

config.OLLAMA_HOST = MOCK_HOST
analyze.config.OLLAMA_HOST = MOCK_HOST

SRC = ROOT / "samples" / "회사보고자료_예시.pptx"
FAILS = 0


def check(label: str, ok: bool, detail: str = "") -> None:
    global FAILS
    print(f"{'PASS' if ok else 'FAIL'}  {label}" + (f"  ({detail})" if detail else ""))
    if not ok:
        FAILS += 1


def runs_of(prs):
    """(슬라이드 번호, 도형 이름, 문단 글자, 런) 을 차례로."""
    for s_no, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs)
                    for r in p.runs:
                        yield s_no, shp.name, t, r


def hl_of(run) -> str | None:
    rPr = run._r.find(qn("a:rPr"))
    c = rPr.find(f"{qn('a:highlight')}/{qn('a:srgbClr')}") if rPr is not None else None
    return c.get("val") if c is not None else None


def is_tool(run) -> bool:
    rPr = run._r.find(qn("a:rPr"))
    uf = rPr.find(f"{qn('a:uFill')}/{qn('a:solidFill')}/{qn('a:srgbClr')}") if rPr is not None else None
    return rPr is not None and rPr.get("u") == "sng" and uf is not None and uf.get("val") in prehl.TOOL_COLORS


def text_color(run) -> str | None:
    rPr = run._r.find(qn("a:rPr"))
    c = rPr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}") if rPr is not None else None
    return c.get("val") if c is not None else None


# ── 0. 기준: 형광펜 없는 원본 ───────────────────────────────────────
base, base_st = pipeline.run(SRC, TMP / "base.pptx", config.RunOptions())
deck0 = extract.extract(str(SRC))
seg0 = {s.seg_id: s for s in deck0.segments}
marked = {seg0[f.seg_id].text: f for f in base if f.span and seg0[f.seg_id].kind != "notes"}
check("형광펜 없는 문서: 보고할 것이 없다", base_st["prehl"]["found"] == 0 and base_st["prehl_text"] == "",
      json.dumps(base_st["prehl"], ensure_ascii=False))

# ── 1. 작성자 형광펜을 넣은 PPTX ────────────────────────────────────
prs = Presentation(str(SRC))
made = {"overlap": None, "plain": [], "dark": None}
for slide in prs.slides:
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        for p in shp.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            if len(t) < 12:
                continue
            f = marked.get(t)
            if f is not None and made["overlap"] is None and f.span[1] - f.span[0] < len(t) - 4:
                made["overlap"] = t                                  # 도구가 일부를 칠하는 문단 전체에 작성자 노랑
                for r in p.runs:
                    mark._set_highlight(r._r, "FFFF00")
            elif f is None and t not in marked and len(made["plain"]) < 2:
                made["plain"].append(t)                              # 도구가 안 잡는 문단: 노랑 · 연두
                for r in p.runs:
                    mark._set_highlight(r._r, "FFFF00" if len(made["plain"]) == 1 else "00FF00")
            elif f is None and t not in marked and made["dark"] is None and not mark._background_is_dark(shp, slide):
                made["dark"] = t                                     # 진한 파랑 형광펜 + 흰 글자 (밝은 바탕)
                for r in p.runs:
                    mark._set_highlight(r._r, "000080")
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
author = TMP / "author.pptx"
prs.save(str(author))
before = author.read_bytes()
check("점검용 문서: 작성자 형광펜 4문단", all((made["overlap"], made["dark"])) and len(made["plain"]) == 2)

out = TMP / "author_marked.pptx"
res, st = pipeline.run(author, out, config.RunOptions())
info = st["prehl"]
check("원본 파일은 바뀌지 않는다", author.read_bytes() == before)
check("형광펜 4곳을 찾아 4곳을 지움", info["mode"] == "strip" and info["found"] == 4 and info["removed"] == 4
      and info["kept"] == 0 and info["tool_prev"] == 0, json.dumps(info, ensure_ascii=False))
check("결과 한 줄", st["prehl_text"].startswith("원본 형광펜 4곳을 지우고 분석") and "글자색 보정 1곳" in st["prehl_text"],
      st["prehl_text"])
check("후보는 형광펜 없는 원본과 같다", sorted((f.slide_no, f.quote, f.grade) for f in res)
      == sorted((f.slide_no, f.quote, f.grade) for f in base), f"{len(res)} vs {len(base)}")

outp = Presentation(str(out))
hl_runs = [(t, r) for _, name, t, r in runs_of(outp) if name != config.LEGEND_NAME and hl_of(r)]
check("마킹본의 형광펜은 전부 도구 표시 (밑줄 서명)", hl_runs and all(is_tool(r) for _, r in hl_runs),
      f"형광펜 런 {len(hl_runs)}개")
check("도구가 안 잡는 문단에는 형광펜이 남지 않는다",
      not any(t in made["plain"] or t == made["dark"] for t, _ in hl_runs))
dark_runs = [r for _, _, t, r in runs_of(outp) if t == made["dark"]]
check("진한 형광펜을 지운 자리의 흰 글자는 어둡게 바뀐다",
      dark_runs and all(text_color(r) == config.MARKED_TEXT_COLOR for r in dark_runs),
      str({text_color(r) for r in dark_runs}))
d = review.diff(review.read_reviewed(str(out)))
check("검토 반영 미리보기(고치지 않음): 가짜 누락·오탐·등급변경 0",
      d["mode"] == "archive" and d["counts"]["miss"] == 0 and d["counts"]["fp"] == 0 and d["counts"]["regrade"] == 0
      and d["counts"]["match"] >= 1, json.dumps(d["counts"], ensure_ascii=False))
log = (config.REVIEW_DIR / "run_log.tsv").read_text(encoding="utf-8-sig").splitlines()
head = next(ln for ln in reversed(log) if ln.startswith("일시\t")).split("\t")
check("run_log.tsv 에 '원본형광펜' 칸", "원본형광펜" in head and log[-1].split("\t")[head.index("원본형광펜")] == "지움 4",
      log[-1].split("\t")[head.index("원본형광펜")] if "원본형광펜" in head else "칸 없음")

# ── 2. 끈 경우: 그대로 두고 경고 ────────────────────────────────────
res_k, st_k = pipeline.run(author, TMP / "kept.pptx", config.RunOptions(strip_highlights=False))
kp = Presentation(str(TMP / "kept.pptx"))
left = [r for _, name, t, r in runs_of(kp) if name != config.LEGEND_NAME and hl_of(r) and not is_tool(r)]
dk = review.diff(review.read_reviewed(str(TMP / "kept.pptx")))
check("끄면 작성자 형광펜이 그대로 남는다", st_k["prehl"]["mode"] == "keep" and st_k["prehl"]["found"] == 4
      and st_k["prehl"]["removed"] == 0 and len(left) >= 3, f"남은 런 {len(left)}개")
check("끄면 경고 한 줄이 나온다", "그대로 두었습니다" in st_k["prehl_text"] and "검토 반영" in st_k["prehl_text"],
      st_k["prehl_text"])
check("끄면 검토 반영에 가짜 누락이 생긴다 (이 기능이 막는 것)", dk["counts"]["miss"] >= 3,
      json.dumps(dk["counts"], ensure_ascii=False))

# ── 3. 마킹본을 다시 분석해도 깨끗하다 ──────────────────────────────
res2, st2 = pipeline.run(out, TMP / "again.pptx", config.RunOptions(tag_marks=True))
ag = Presentation(str(TMP / "again.pptx"))
legends = [s for sl in ag.slides for s in sl.shapes if s.name == config.LEGEND_NAME]
check("예전 도구 표시를 걷어낸다 (형광펜·범례)", st2["prehl"]["tool_prev"] >= len(base) and st2["prehl"]["found"] == 0,
      json.dumps(st2["prehl"], ensure_ascii=False))
check("범례 상자는 하나만", len(legends) == 1, f"{len(legends)}개")
check("후보 수가 같다", len(res2) == len(base), f"{len(res2)} vs {len(base)}")
res3, st3 = pipeline.run(TMP / "again.pptx", TMP / "again2.pptx", config.RunOptions())
texts3 = ["".join(r.text for r in p.runs) for sl in Presentation(str(TMP / "again2.pptx")).slides
          for shp in sl.shapes if shp.has_text_frame for p in shp.text_frame.paragraphs]
check("예전에 붙인 【출원검토필요】 문구도 걷어낸다", not any(config.TAG_TEXT in t for t in texts3) and len(res3) == len(base),
      f"후보 {len(res3)}")
check("지문이 원본과 같다 (분석 기록이 이어진다)",
      review.read_reviewed(str(TMP / "again2.pptx"))["fingerprint"] == review.deck_fingerprint(deck0))

# ── 4. PDF: 뷰어로 칠한 형광펜 주석 ─────────────────────────────────
pdf = TMP / "src.pdf"
make_sample_pdf.build(SRC, pdf)
pbase, _ = pipeline.run(pdf, TMP / "pbase.pdf", config.RunOptions())
pdeck = extract.extract(str(pdf))
pmarked = {f.seg_id for f in pbase}
w = PdfWriter(clone_from=PdfReader(str(pdf)))
um = next(s for s in pdeck.segments if s.seg_id not in pmarked and len(s.text) >= 15)
mk = next(s for s in pdeck.segments if s.seg_id in pmarked and len(s.text) >= 15)
for seg in (um, mk):
    rects = pdfdoc._line_rects(pdeck.boxes[seg.seg_id], None)
    w.add_annotation(page_number=seg.slide_no - 1, annotation=Highlight(
        rect=pdfdoc._union(rects), quad_points=pdfdoc._quads(rects), highlight_color="FFFF00", title_bar="작성자"))
w.add_annotation(page_number=0, annotation=Text(rect=(300, 300, 320, 320), text="회의 메모", title_bar="작성자"))
apdf = TMP / "author.pdf"
with open(apdf, "wb") as fh:
    w.write(fh)
apdf_before = apdf.read_bytes()


def annots(path):
    return [a.get_object() for pg in PdfReader(str(path)).pages for a in (pg.get("/Annots") or [])]


pres, pst = pipeline.run(apdf, TMP / "author_marked.pdf", config.RunOptions())
an = annots(TMP / "author_marked.pdf")
hls = [a for a in an if a.get("/Subtype") == "/Highlight"]
check("PDF 원본은 바뀌지 않는다", apdf.read_bytes() == apdf_before)
check("작성자 형광펜 주석 2개를 지운다", pst["prehl"]["found"] == 2 and pst["prehl"]["removed"] == 2,
      json.dumps(pst["prehl"], ensure_ascii=False))
check("남은 형광펜 주석은 전부 도구의 것", len(hls) == len(pres) and all(str(a.get("/T")) == "출원검토필요" for a in hls),
      f"주석 {len(hls)} · 후보 {len(pres)}")
check("형광펜이 아닌 주석(메모)은 그대로", any(a.get("/Subtype") == "/Text" and str(a.get("/T")) == "작성자" for a in an))
pdv = review.diff(review.read_reviewed(str(TMP / "author_marked.pdf")))
check("PDF 검토 반영 미리보기: 가짜 항목 0", pdv["counts"]["miss"] == 0 and pdv["counts"]["regrade"] == 0
      and pdv["counts"]["fp"] == 0 and pdv["counts"]["match"] == len(pres), json.dumps(pdv["counts"], ensure_ascii=False))
pk, pkst = pipeline.run(apdf, TMP / "author_kept.pdf", config.RunOptions(strip_highlights=False))
pkd = review.diff(review.read_reviewed(str(TMP / "author_kept.pdf")))
check("PDF 도 끄면 남고, 가짜 항목이 생긴다", pkst["prehl"]["mode"] == "keep" and pkst["prehl"]["found"] == 2
      and pkd["counts"]["miss"] + pkd["counts"]["regrade"] >= 2, json.dumps(pkd["counts"], ensure_ascii=False))
pa, past = pipeline.run(TMP / "author_marked.pdf", TMP / "author_again.pdf", config.RunOptions())
an2 = annots(TMP / "author_again.pdf")
check("PDF 마킹본 재분석: 예전 도구 주석을 걷어내고 범례는 한 벌만",
      past["prehl"]["tool_prev"] == len(pres) + 4 and past["prehl"]["found"] == 0
      and sum(a.get("/Subtype") == "/Highlight" for a in an2) == len(pa)
      and sum(str(a.get("/T")) == config.LEGEND_NAME for a in an2) == 4, json.dumps(past["prehl"], ensure_ascii=False))

# ── 5. PDF: 쪽에 '그림'으로 굳은 형광펜 ─────────────────────────────
p2 = [s for s in pdeck.segments if s.slide_no == 2]
long_ = [i for i, s in enumerate(p2, 1) if len(s.text) >= 30]              # 긴 문단 셋: 노랑 · 청록 · 노랑+흰 글자
short_ = [i for i, s in enumerate(p2, 1) if 5 <= len(s.text) < 30]
HL = [{"slide": 2, "para": long_[0], "color": "FFFF00"},
      {"slide": 2, "para": long_[1], "color": "00FFFF"},
      {"slide": 2, "para": long_[2], "color": "FFFF00", "text": "FFFFFF"},       # 밝은 글자 → 그대로 둠
      {"slide": 2, "para": short_[0], "color": "000080", "text": "FFFFFF"},      # 진한 형광펜 → 대상 아님
      {"slide": 2, "para": short_[1], "color": "66B2FF", "pad": 8}]              # 여백 둔 채움(표 칸·도형) → 형광펜 아님
bpdf = TMP / "baked.pdf"
make_sample_pdf.build(SRC, bpdf, highlights=HL)
bdeck = pdfdoc.extract(str(bpdf))
found = [(pdfdoc._hex(c["rgb"]), pdfdoc._hex(c["bg"])) for c in bdeck.baked.get(2, [])]
check("굳은 형광펜 2곳을 찾는다 (색을 제대로 읽는다: 저장·복원 뒤에도)",
      sorted(found) == [("00FFFF", "FFFFFF"), ("FFFF00", "FFFFFF")] and set(bdeck.baked) == {2}, str(found))
check("밝은 글자 1곳은 그대로 둘 것으로 센다 (진한 형광펜·여백 둔 채움은 세지 않는다)", bdeck.baked_kept == 1,
      f"kept {bdeck.baked_kept}")
check("글자는 형광펜 없는 PDF 와 똑같이 읽힌다", [s.text for s in bdeck.segments] == [s.text for s in pdeck.segments])

bres, bst = pipeline.run(bpdf, TMP / "baked_marked.pdf", config.RunOptions())
rd = PdfReader(str(TMP / "baked_marked.pdf"))
page2 = rd.pages[1]
content = page2.get_contents().get_data().decode("latin-1")
ext = page2["/Resources"]["/ExtGState"]
modes = {k: str(ext[k]["/BM"]) for k in ext if k.startswith("/PMhl")}
check("덮기: 3곳 중 2곳을 덮고 1곳은 둠", bst["prehl"]["found"] == 3 and bst["prehl"]["covered"] == 2
      and bst["prehl"]["kept"] == 1 and bst["prehl"]["removed"] == 0, json.dumps(bst["prehl"], ensure_ascii=False))
check("덮기 그림이 2쪽에만 들어간다", "/PMhlS gs" in content and "/PMhlD gs" in content
      and "/PMhlS gs" not in rd.pages[0].get_contents().get_data().decode("latin-1"))
check("혼합 모드 넷이 등록된다", modes == {"/PMhlS": "/Saturation", "/PMhlD": "/ColorDodge",
                                       "/PMhlM": "/Multiply", "/PMhlC": "/Color"}, str(modes))
check("덮은 뒤에도 글자·쪽 수가 같다", len(rd.pages) == 5
      and [s.text for s in pdfdoc.extract(str(TMP / "baked_marked.pdf"), detect_baked=False).segments]
      == [s.text for s in pdeck.segments])
check("덮은 PDF 의 검토 반영 미리보기: 가짜 항목 0",
      review.diff(review.read_reviewed(str(TMP / "baked_marked.pdf")))["counts"]["miss"] == 0)
check("결과 한 줄 (덮음·그대로 둠)", "3곳 중 2곳" in bst["prehl_text"] and "그대로 둠" in bst["prehl_text"]
      and "덮음" in bst["prehl_text"], bst["prehl_text"])
for _t in (st["prehl_text"], st_k["prehl_text"], bst["prehl_text"], st2["prehl_text"]):
    try:
        _t.encode("cp949")
        _ok = True
    except UnicodeEncodeError:
        _ok = False
    check("결과 한 줄은 한국어 Windows 콘솔 글자표(cp949)로 찍을 수 있다", _ok, _t[:50])
ops = prehl._cover_ops(bdeck.baked[2])
check("덮는 세기는 형광펜 밝기에 맞춘다 (노랑 0.14 · 청록 0.33 — 글자가 가늘어지지 않게)",
      "0.1400 g" in ops and "0.3300 g" in ops, ops.replace("\n", " ")[:160])

config.PDF_COVER_BAKED = False
try:
    nres, nst = pipeline.run(bpdf, TMP / "baked_nocover.pdf", config.RunOptions())
finally:
    config.PDF_COVER_BAKED = True
check("덮기만 끌 수 있다 (PM_PDF_COVER=0): 3곳 모두 그대로",
      nst["prehl"]["covered"] == 0 and nst["prehl"]["kept"] == 3
      and "/PMhlS gs" not in PdfReader(str(TMP / "baked_nocover.pdf")).pages[1].get_contents().get_data().decode("latin-1"),
      json.dumps(nst["prehl"], ensure_ascii=False))

# ── 6. 한 줄 문구 · 옵션 전달 ───────────────────────────────────────
check("문구: 없으면 빈 줄 · 실행 기록 칸", prehl.text({"mode": "strip", "found": 0}) == "" and prehl.short(None) == "0"
      and prehl.short({"mode": "keep", "found": 3}) == "남김 3"
      and prehl.short({"mode": "strip", "found": 3, "removed": 0, "covered": 2, "kept": 1}) == "지움 2·남김 1")
env = {**os.environ, "PYTHONIOENCODING": "utf-8",         # 자식의 출력 글자표 고정 (Windows 기본은 cp949)
       "PM_OLLAMA_HOST": MOCK_HOST, "PM_OUTPUT_DIR": str(config.OUTPUT_DIR), "PM_REVIEW_DIR": str(config.REVIEW_DIR)}
proc = subprocess.run([sys.executable, "-m", "app.cli", str(author), "--out", str(TMP / "cli_keep"), "--keep-highlights"],
                      cwd=str(ROOT), env=env, capture_output=True, text=True, encoding="utf-8")
cli_out = TMP / "cli_keep" / "author_특허마킹.pptx"
cli_left = [r for _, name, t, r in runs_of(Presentation(str(cli_out))) if name != config.LEGEND_NAME and hl_of(r) and not is_tool(r)] \
    if cli_out.exists() else []
check("명령행 --keep-highlights: 남기고, 검은 창에 알린다", proc.returncode == 0 and len(cli_left) >= 3
      and "원본 형광펜" in proc.stdout and "그대로 두었습니다" in proc.stdout, (proc.stderr or "")[-160:])
proc2 = subprocess.run([sys.executable, "-m", "app.cli", str(author), "--out", str(TMP / "cli_strip")],
                       cwd=str(ROOT), env=env, capture_output=True, text=True, encoding="utf-8")
check("명령행 기본: 지우고 알린다", proc2.returncode == 0 and "원본 형광펜 4곳을 지우고 분석" in proc2.stdout,
      (proc2.stderr or proc2.stdout)[-160:] if proc2.returncode else "")

from app import main as webmain  # noqa: E402

check("화면 옵션 → RunOptions (strip_hl)", webmain._options(config.MODEL, False, True, False, True, True, False).strip_highlights is False
      and webmain._options(config.MODEL, False, True, False).strip_highlights is True)
page = (ROOT / "app" / "static" / "index.html").read_text(encoding="utf-8")
check("화면에 체크박스와 결과 줄이 있다", 'id="optStrip" checked' in page and "strip_hl" in page and "st.prehl_text" in page)

print(f"\n{'모두 통과' if not FAILS else f'실패 {FAILS}건'}  (임시 폴더: {TMP})")
sys.exit(1 if FAILS else 0)
