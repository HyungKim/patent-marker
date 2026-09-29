"""
tests/test_cp949.py ─ 한국어 Windows 의 검은 창 조건 흉내 (LLM 없이, 30초)
=====================================================================
한국어 Windows 에서는 프로그램의 출력이 화면이 아닌 곳(파일·다른 프로그램)으로 넘어가면
글자표가 cp949 가 됩니다. cp949 에는 — ✗ • 같은 글자가 없어서, 그런 글자를 찍는 순간
프로그램이 죽을 수 있습니다. 이 점검은 출력 글자표를 cp949 로 고정한 채
  · 명령행(app/cli.py = mark.bat 이 부르는 것)으로 PPTX·PDF 를 끝까지 마킹하고
  · 웹 서버(app/main.py = run.bat 이 부르는 것)를 띄워 분석을 한 번 돌려
죽지 않는지, 찍힌 글이 cp949 로 읽히는지 확인합니다. Mac·Linux 에서도 같은 조건으로 돌아갑니다.

    실행:  python tests/test_cp949.py
"""
from __future__ import annotations

import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:  # noqa: BLE001
        pass

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
sys.path.insert(0, str(ROOT / "tools"))

TMP = Path(tempfile.mkdtemp(prefix="pm-cp949-test-"))
os.environ["PM_REVIEW_DIR"] = str(TMP / "review_data")
os.environ["PM_INPUT_DIR"] = str(TMP / "input")
os.environ["PM_OUTPUT_DIR"] = str(TMP / "output")

import mock_llm  # noqa: E402

MOCK_HOST = mock_llm.start()

from pptx import Presentation  # noqa: E402

import make_sample_pdf  # noqa: E402
from app import config, mark  # noqa: E402

SRC = ROOT / "samples" / "회사보고자료_예시.pptx"
FAILS = 0


def check(label: str, ok: bool, detail: str = "") -> None:
    global FAILS
    print(f"{'PASS' if ok else 'FAIL'}  {label}" + (f"  ({detail})" if detail else ""))
    if not ok:
        FAILS += 1


def decodes(b: bytes) -> bool:
    try:
        b.decode("cp949")
        return True
    except UnicodeDecodeError:
        return False


# 작성자 형광펜이 든 문서 — 검은 창에 "원본 형광펜 …" 줄까지 찍히게
prs = Presentation(str(SRC))
n = 0
for slide in prs.slides:
    for shp in slide.shapes:
        if shp.has_text_frame:
            for p in shp.text_frame.paragraphs:
                if n < 3 and len("".join(r.text for r in p.runs)) >= 12:
                    n += 1
                    for r in p.runs:
                        mark._set_highlight(r._r, "FFFF00")
folder = TMP / "보고 자료 (공백·한글 경로)"
folder.mkdir()
pptx = folder / "형광펜 든 보고서.pptx"
prs.save(str(pptx))
pdf = folder / "변환본.pdf"
make_sample_pdf.build(SRC, pdf, highlights=[{"slide": 2, "para": 5, "color": "FFFF00"},
                                            {"slide": 2, "para": 8, "color": "FFFF00", "text": "FFFFFF"}])

ENV = {**os.environ, "PYTHONIOENCODING": "cp949",          # ← 한국어 Windows 에서 출력을 넘길 때의 글자표
       "PM_OLLAMA_HOST": MOCK_HOST, "PM_OUTPUT_DIR": str(TMP / "output"),
       "PM_REVIEW_DIR": str(TMP / "review_data"), "PM_INPUT_DIR": str(TMP / "input")}
ENV.pop("PYTHONUTF8", None)

# ── 1. 명령행 ────────────────────────────────────────────────────
proc = subprocess.run([sys.executable, "-m", "app.cli", str(pptx), str(pdf), "--out", str(TMP / "결과")],
                      cwd=str(ROOT), env=ENV, capture_output=True)
out = proc.stdout.decode("cp949", "replace")
check("명령행: 끝까지 돌고 0 으로 끝난다", proc.returncode == 0, proc.stderr.decode("cp949", "replace")[-200:])
check("명령행: 찍힌 글이 cp949 로 읽힌다", decodes(proc.stdout) and decodes(proc.stderr))
check("명령행: 글자표 오류로 죽은 흔적이 없다", b"UnicodeEncodeError" not in proc.stderr and b"Traceback" not in proc.stderr)
check("명령행: 결과 두 개가 저장된다 (공백·한글 경로)",
      (TMP / "결과" / "형광펜 든 보고서_특허마킹.pptx").exists() and (TMP / "결과" / "변환본_특허마킹.pdf").exists())
check("명령행: 원본 형광펜 줄이 한글로 찍힌다", "원본 형광펜 3곳을 지우고 분석" in out and "그대로 둠" in out,
      next((ln.strip() for ln in out.splitlines() if "원본 형광펜" in ln), "없음"))

proc = subprocess.run([sys.executable, "-m", "app.cli", str(pptx), "--out", str(TMP / "결과2"), "--keep-highlights"],
                      cwd=str(ROOT), env=ENV, capture_output=True)
check("명령행 --keep-highlights: 경고 줄도 cp949 로 찍힌다", proc.returncode == 0 and decodes(proc.stdout)
      and "그대로 두었습니다" in proc.stdout.decode("cp949", "replace"))

proc = subprocess.run([sys.executable, "-m", "app.cli", str(folder / "없는 파일.pptx")],
                      cwd=str(ROOT), env=ENV, capture_output=True)
check("명령행: 없는 파일은 안내만 하고 끝난다 (✗ 는 cp949 에 없어 ? 로 바뀜, 죽지 않음)",
      proc.returncode != 0 and b"UnicodeEncodeError" not in proc.stderr and decodes(proc.stdout),
      proc.stdout.decode("cp949", "replace").strip()[:60])

# ── 2. 웹 서버 ───────────────────────────────────────────────────
with socket.socket() as sk:
    sk.bind(("127.0.0.1", 0))
    port = sk.getsockname()[1]
log = open(TMP / "server.log", "wb")
srv = subprocess.Popen([sys.executable, "-u", "-m", "app.main"], cwd=str(ROOT), env={**ENV, "PM_PORT": str(port)},
                       stdout=log, stderr=subprocess.STDOUT)
BASE = f"http://127.0.0.1:{port}"
try:
    up = False
    for _ in range(150):
        if srv.poll() is not None:
            break
        try:
            with urllib.request.urlopen(BASE + "/api/health", timeout=2) as r:
                health = json.loads(r.read())
                up = True
                break
        except Exception:  # noqa: BLE001
            time.sleep(0.2)
    check("서버: cp949 출력으로도 뜬다", up and srv.poll() is None)
    page = urllib.request.urlopen(BASE + "/", timeout=10).read().decode("utf-8") if up else ""
    token = page.split('const TOKEN = "', 1)[1].split('"', 1)[0] if 'const TOKEN = "' in page else ""
    check("서버: 첫 화면에 버전과 토큰", config.VERSION in page and len(token) > 10)
    (TMP / "input").mkdir(exist_ok=True)
    req = urllib.request.Request(BASE + "/api/local/jobs", method="POST",
                                 data=json.dumps({"path": str(pptx), "model": config.MODEL}).encode("utf-8"),
                                 headers={"Content-Type": "application/json", "X-PM-Token": token})
    job = json.loads(urllib.request.urlopen(req, timeout=30).read())["job_id"] if up else ""
    snap = {}
    for _ in range(300):
        with urllib.request.urlopen(f"{BASE}/api/jobs/{job}", timeout=10) as r:
            snap = json.loads(r.read())
        if snap.get("status") in ("done", "error"):
            break
        time.sleep(0.2)
    check("서버: 경로로 넣은 파일을 끝까지 분석한다", snap.get("status") == "done", str(snap.get("error") or snap.get("stage")))
    summ = snap.get("summary") or snap.get("stats") or {}
    check("서버: 원본 형광펜 줄이 결과에 담긴다", "원본 형광펜 3곳을 지우고 분석" in json.dumps(snap, ensure_ascii=False),
          str(summ.get("prehl_text", ""))[:60])
finally:
    srv.terminate()
    try:
        srv.wait(timeout=10)
    except subprocess.TimeoutExpired:
        srv.kill()
    log.close()
raw = (TMP / "server.log").read_bytes()
check("서버: 검은 창 첫 줄이 cp949 로 찍히고, 글자표 오류가 없다",
      decodes(raw) and b"UnicodeEncodeError" not in raw and "특허 마킹 도구" in raw.decode("cp949", "replace"),
      raw.decode("cp949", "replace").strip().splitlines()[0][:60] if raw.strip() else "출력 없음")

print(f"\n{'모두 통과' if not FAILS else f'실패 {FAILS}건'}  (임시 폴더: {TMP})")
sys.exit(1 if FAILS else 0)
