"""
prehl.py ─ 원본에 있던 형광펜 정리 ("지우고 시작하기")
=====================================================================

[왜 필요한가]
  보고서에는 작성자가 강조하려고 칠해 둔 형광펜이 들어 있는 경우가 많습니다. 그대로 두면
    · 마킹본에서 도구의 표시와 섞여 보이고 (작성자 노랑 = 도구의 A 표시와 같은 색)
    · 검토 반영 때 "검토자가 새로 칠한 누락" 으로 잘못 읽혀, 엉뚱한 표현이 확정 후보로 학습됩니다.
  그래서 분석을 시작할 때 원본의 형광펜을 걷어냅니다. **원본 파일은 건드리지 않습니다** —
  걷어내는 것은 메모리에 올린 사본(= 저장되는 마킹본)에서만입니다.
  그 결과 마킹본에 남는 형광펜은 전부 "도구가 칠한 것" 또는 "검토자가 고친 것" 이 됩니다.

[하는 일]
  PPTX   clean_pptx(prs, strip)   글자(런)의 형광펜 제거 + 글자색 보정 + 도구가 예전에 넣은 표시(밑줄·문구·범례) 제거
  PDF    find_baked(layout)       쪽 내용에 '그림'으로 굳어 있는 형광펜(글줄에 딱 붙은 채움 네모) 찾기
         clean_pdf(writer, ...)   형광펜 주석 삭제 + 굳은 형광펜을 바탕색으로 덮기
  text(info)                      화면·검은 창·실행 기록에 쓰는 한 줄

[PDF 의 두 가지 형광펜]
  ① 주석 형광펜 : PDF 뷰어(Edge·Acrobat)에서 칠한 것. 파일 안에 '주석'으로 들어 있어 그냥 지우면 됩니다.
  ② 굳은 형광펜 : 파워포인트를 PDF 로 저장할 때 색 네모 '그림'으로 들어간 것. 지울 수 없어서, 그 위에
     "색을 빼고(채도 혼합) → 바탕 밝기로 올리는(닷지 혼합)" 칠을 덧그려 안 보이게 합니다. 글자는 그대로 남습니다.
     검토 반영은 주석만 읽으므로 ②는 학습에는 영향이 없고, 눈에 섞여 보이는 것을 막는 용도입니다.
     확신이 없는 네모(표 칸 채우기·도형·차트 막대, 밝은 글자, 어두운 바탕)는 건드리지 않습니다.

[끄고 싶다면]  화면의 "원본 형광펜 지우고 시작" 체크를 풀거나, 명령행 --keep-highlights, 환경 변수 PM_STRIP_HL=0.
               굳은 형광펜 덮기만 끄려면 PM_PDF_COVER=0 (또는 config.PDF_COVER_BAKED = False).
"""
from __future__ import annotations

import colorsys

from pptx.oxml.ns import qn

from . import config

# 도구가 칠하는 색 (지금 + 예전 기본색). 밑줄 채움색이 이 중 하나면 "도구가 예전에 칠한 표시" 로 본다.
TOOL_COLORS = frozenset(c.upper() for c in (*config.GRADE_COLOR.values(), *config.LEGACY_GRADE_COLORS))
TOOL_SHAPE_NAMES = frozenset((config.LEGEND_NAME, *config.LEGACY_BADGE_NAMES))
TAG_TEXTS = frozenset((config.TAG_TEXT, *config.LEGACY_TAG_TEXTS))
TOOL_ANNOT_TITLES = frozenset(("출원검토필요", config.LEGEND_NAME))


def new_info(strip: bool) -> dict:
    return {"mode": "strip" if strip else "keep",
            "found": 0,          # 원본에 있던 형광펜 구간 수 (도구가 예전에 칠한 것은 빼고)
            "removed": 0,        # 지운 구간 (PPTX 런 형광펜 · PDF 주석)
            "covered": 0,        # 덮은 구간 (PDF 굳은 형광펜)
            "kept": 0,           # 확신이 없어 그대로 둔 구간
            "text_fixed": 0,     # 형광펜을 지우면 안 보일 글자를 보정한 구간
            "tool_prev": 0}      # 도구가 예전에 넣은 표시(형광펜·범례·문구)를 걷어낸 수


def text(info: dict | None) -> str:
    """결과 한 줄. 원본 형광펜이 없었으면 빈 문자열.
    (검은 창에도 찍히는 글이라, 한국어 Windows 콘솔 글자표(cp949)에 없는 줄표 — 는 쓰지 않는다)"""
    if not info:
        return ""
    found, prev = info.get("found", 0), info.get("tool_prev", 0)
    if not found and not prev:
        return ""
    if info.get("mode") != "strip":
        return (f"원본 형광펜 {found}곳을 그대로 두었습니다. 검토 반영 때 검토자가 칠한 것으로 읽힐 수 있습니다"
                if found else "")
    bits = []
    if found:
        done, kept = info.get("removed", 0) + info.get("covered", 0), info.get("kept", 0)
        s = (f"원본 형광펜 {found}곳 중 {done}곳을 지우고 분석" if kept else f"원본 형광펜 {done}곳을 지우고 분석")
        extra = []
        if kept:
            extra.append(f"{kept}곳은 밝은 글자이거나 어두운 바탕이라 그대로 둠")
        if info.get("covered"):
            extra.append(f"PDF 에 굳어 있던 {info['covered']}곳은 바탕색으로 덮음")
        if info.get("text_fixed"):
            extra.append(f"글자색 보정 {info['text_fixed']}곳")
        bits.append(s + (f" ({' · '.join(extra)})" if extra else ""))
    if prev:
        bits.append(f"예전 도구 표시 {prev}개 제거")
    return " · ".join(bits)


def short(info: dict | None) -> str:
    """실행 기록(run_log.tsv) 칸에 적는 짧은 값. 예: '지움 7' · '지움 5·남김 1' · '남김 7' · '0'"""
    if not info or not info.get("found"):
        return "0"
    if info.get("mode") != "strip":
        return f"남김 {info['found']}"
    done = info.get("removed", 0) + info.get("covered", 0)
    return f"지움 {done}" + (f"·남김 {info['kept']}" if info.get("kept") else "")


# ═════════════════════════════════════════════════════════════════
# 1. PPTX
# ═════════════════════════════════════════════════════════════════
def _walk(shapes):
    for shp in shapes:
        if shp.shape_type == 6 and hasattr(shp, "shapes"):      # 6 = GROUP
            yield from _walk(shp.shapes)
        else:
            yield shp


def _cell_is_dark(cell, mark) -> bool | None:
    """표 칸의 채우기 색으로 본 바탕 밝기. 채우기가 없거나 테마 색이면 None."""
    tcPr = cell._tc.find(qn("a:tcPr"))
    lum = mark._luma(mark._direct_fill_hex(tcPr)) if tcPr is not None else None
    return None if lum is None else lum < 0.45


def _clean_paragraph(p_el, strip: bool, dark_bg: bool, info: dict, mark) -> None:
    prev_key = None                       # 바로 앞 런의 (형광펜 색, 도구 표시 여부) — 이어지면 같은 구간으로 센다
    for run in list(p_el):
        if run.tag not in (qn("a:r"), qn("a:fld")):
            continue
        t = run.find(qn("a:t"))
        if strip and t is not None and (t.text or "").strip() in TAG_TEXTS:
            p_el.remove(run)              # 도구가 예전에 붙인 【출원검토필요】 문구
            info["tool_prev"] += 1
            prev_key = None
            continue
        rPr = run.find(qn("a:rPr"))
        hl = rPr.find(qn("a:highlight")) if rPr is not None else None
        if hl is None:
            prev_key = None
            continue
        clr = hl.find(qn("a:srgbClr"))
        color = (clr.get("val") or "").upper() if clr is not None else "?"
        ufill = rPr.find(f"{qn('a:uFill')}/{qn('a:solidFill')}/{qn('a:srgbClr')}")
        is_tool = rPr.get("u") == "sng" and ufill is not None and (ufill.get("val") or "").upper() in TOOL_COLORS
        key = (color, is_tool)
        new_span = key != prev_key
        prev_key = key
        if new_span:
            info["tool_prev" if is_tool else "found"] += 1
        if not strip:
            continue
        rPr.remove(hl)
        if is_tool:                       # 도구의 서명(같은 색 밑줄)도 함께 걷어낸다
            rPr.attrib.pop("u", None)
            for uf in rPr.findall(qn("a:uFill")):
                rPr.remove(uf)
            continue
        if new_span:
            info["removed"] += 1
        # 글자색 보정 — 형광펜에 맞춰 고른 글자색이라 형광펜이 없어지면 안 보이게 되는 경우만
        hl_l = mark._luma(color) if color != "?" else None
        own = rPr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
        t_l = mark._luma(own.get("val")) if own is not None else None
        if hl_l is None or t_l is None:
            continue
        if hl_l < 0.45 and t_l > 0.6 and not dark_bg:          # 진한 형광펜 + 밝은 글자, 바탕은 밝음
            mark._set_text_color(run, config.MARKED_TEXT_COLOR)
            info["text_fixed"] += 1 if new_span else 0
        elif hl_l >= 0.45 and t_l < 0.4 and dark_bg:           # 밝은 형광펜 + 어두운 글자, 바탕은 어두움
            mark._set_text_color(run, "FFFFFF")
            info["text_fixed"] += 1 if new_span else 0


def clean_pptx(prs, strip: bool) -> dict:
    """문서(메모리 사본)의 형광펜을 세고, strip 이면 걷어낸다. 글자와 런 경계는 바꾸지 않는다
    (예전 도구 문구 런을 지우는 경우만 예외 — 그 문구는 원문이 아니다)."""
    from . import mark                   # 늦게 불러 순환 참조를 피한다 (mark → extract → prehl)

    info = new_info(strip)
    for slide in prs.slides:
        if strip:                         # 도구가 예전에 붙인 범례 상자·(옛)배지
            for shp in list(slide.shapes):
                if getattr(shp, "name", "") in TOOL_SHAPE_NAMES:
                    shp._element.getparent().remove(shp._element)
                    info["tool_prev"] += 1
        for shp in _walk(slide.shapes):
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    for cell in row.cells:
                        dark = _cell_is_dark(cell, mark)
                        if dark is None:
                            dark = mark._background_is_dark(None, slide)
                        for para in cell.text_frame.paragraphs:
                            _clean_paragraph(para._p, strip, dark, info, mark)
            elif getattr(shp, "has_text_frame", False) and shp.has_text_frame:
                dark = mark._background_is_dark(shp, slide)
                for para in shp.text_frame.paragraphs:
                    _clean_paragraph(para._p, strip, dark, info, mark)
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                for para in tf.paragraphs:
                    _clean_paragraph(para._p, strip, False, info, mark)
        if strip:                         # 문단 기본값·문단 끝 서식 등에 남은 형광펜 (보이는 글자에는 영향 없음)
            for hl in list(slide._element.iter(qn("a:highlight"))):
                hl.getparent().remove(hl)
    return info


# ═════════════════════════════════════════════════════════════════
# 2. PDF ─ 굳은 형광펜 찾기
# ═════════════════════════════════════════════════════════════════
def _rgb(color) -> tuple[float, float, float] | None:
    """pdfminer 의 색 값 → RGB. 회색(숫자 하나)·RGB·CMYK 를 받는다. 무늬·이름 색 등은 None."""
    try:
        if isinstance(color, (int, float)):
            v = float(color)
            return (v, v, v)
        vals = [float(v) for v in color]
    except (TypeError, ValueError):
        return None
    if len(vals) == 1:
        return (vals[0],) * 3
    if len(vals) == 3:
        return tuple(min(1.0, max(0.0, v)) for v in vals)
    if len(vals) == 4:
        c, m, y, k = vals
        return ((1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k))
    return None


def _lum(rgb) -> float:
    """PDF 혼합 모드가 쓰는 밝기 공식 (채도 혼합의 결과 회색 값과 같다)."""
    return 0.30 * rgb[0] + 0.59 * rgb[1] + 0.11 * rgb[2]


def _is_rect(obj) -> bool:
    pts = getattr(obj, "pts", None) or []
    if len(pts) not in (4, 5):
        return False
    x0, y0, x1, y1 = obj.bbox
    return all(min(abs(px - x0), abs(px - x1)) < 0.2 and min(abs(py - y0), abs(py - y1)) < 0.2 for px, py in pts)


def find_baked(layout) -> tuple[list[dict], int]:
    """쪽 하나(pdfminer LTPage)에서 굳은 형광펜을 찾는다. (덮을 네모 목록, 그대로 둘 개수) 를 돌려준다.

    형광펜으로 보는 조건 (전부 만족해야 덮는다 — 하나라도 어긋나면 건드리지 않는다):
      · 색이 밝고 선명한 채움 네모
      · 글줄 하나에 딱 붙어 있다: 높이가 글줄의 0.8~1.7배, 글줄이 네모 안에 세로로 들어오고,
        좌우로는 글줄 범위를 벗어나지 않는다 (표 칸·도형 채우기는 글자보다 넓게 여백을 둔다)
      · 그 안의 글자가 형광펜보다 충분히 어둡다 (밝은 글자는 덮으면 흐려지거나 안 보이게 된다)
      · 그 자리의 바탕이 밝다 (어두운 바탕은 흰색으로 덮을 수 없다)
    """
    from pdfminer.layout import LTChar, LTCurve, LTFigure, LTTextBox, LTTextLine

    def walk(o):
        for x in o:
            yield x
            if isinstance(x, LTFigure):
                yield from walk(x)

    objs = list(walk(layout))
    lines = [ln for x in objs if isinstance(x, LTTextBox) for ln in x if isinstance(ln, LTTextLine)]
    chars = [c for ln in lines for c in ln if isinstance(c, LTChar)]
    page_area = max(1.0, float(layout.width) * float(layout.height))

    fills = []                            # 그려진 순서대로: (순번, 네모, RGB)
    for i, x in enumerate(objs):
        if isinstance(x, LTCurve) and getattr(x, "fill", False) and _is_rect(x):
            rgb = _rgb(getattr(x, "non_stroking_color", None))
            if rgb is not None and x.width > 0.5 and x.height > 0.5:
                fills.append((i, tuple(float(v) for v in x.bbox), rgb))

    # 1) 형광펜 색 후보만 골라, 같은 줄에서 맞닿은 같은 색 네모를 하나로 잇는다 (파워포인트는 낱말마다 따로 그린다)
    cands = []
    for i, box, rgb in fills:
        h, s, v = colorsys.rgb_to_hsv(*rgb)
        if s >= 0.25 and v >= 0.7 and 4.0 <= (box[3] - box[1]) <= 90.0:
            cands.append([i, list(box), rgb])
    cands.sort(key=lambda c: (round(c[1][1], 0), c[1][0]))
    merged: list[list] = []
    for c in cands:
        m = merged[-1] if merged else None
        if (m is not None and m[2] == c[2] and abs(m[1][1] - c[1][1]) <= 0.6 and abs(m[1][3] - c[1][3]) <= 0.6
                and -0.5 <= c[1][0] - m[1][2] <= 1.0):
            m[1][2] = c[1][2]
            m[0] = min(m[0], c[0])
        else:
            merged.append([c[0], list(c[1]), c[2]])

    cover, kept = [], 0
    for idx, (x0, y0, x1, y1), rgb in merged:
        hh = y1 - y0
        line = None
        for ln in lines:
            lh = ln.height
            if lh <= 0 or not (0.8 * lh <= hh <= 1.7 * lh):
                continue
            ov = min(y1, ln.y1) - max(y0, ln.y0)
            if ov >= 0.6 * lh and x0 >= ln.x0 - 3.0 and x1 <= ln.x1 + 3.0:
                line = ln
                break
        if line is None:
            continue                      # 글줄에 붙지 않은 네모 = 표·도형·차트 → 형광펜이 아님
        inside = [c for c in chars if x0 <= (c.x0 + c.x1) / 2 <= x1 and y0 <= (c.y0 + c.y1) / 2 <= y1]
        if not any(c.get_text().strip() for c in inside):
            continue                      # 글자가 하나도 없는 자리
        lums = []
        for c in inside:
            crgb = _rgb(getattr(getattr(c, "graphicstate", None), "ncolor", None))
            if crgb is not None and c.get_text().strip():
                lums.append(_lum(crgb))
        if lums and max(lums) > 0.4 * _lum(rgb):
            kept += 1                     # 글자가 형광펜에 비해 충분히 어둡지 않다 → 덮으면 글자가 흐려지거나 안 보인다
            continue
        # 이 자리의 바탕: 먼저 그려진 채움 네모 중 이 네모를 넉넉히 감싸는 마지막 것 (없으면 흰 종이)
        bg = (1.0, 1.0, 1.0)
        area = (x1 - x0) * hh
        for j, b, brgb in fills:
            if j >= idx:
                break
            if b[0] <= x0 + 0.5 and b[1] <= y0 + 0.5 and b[2] >= x1 - 0.5 and b[3] >= y1 - 0.5 \
                    and (b[2] - b[0]) * (b[3] - b[1]) >= 3.0 * area:
                bg = brgb
        if _lum(bg) < 0.75:
            kept += 1                     # 어두운 바탕 위의 형광펜 → 그대로 둔다
            continue
        cover.append({"rect": (x0, y0, x1, y1), "rgb": rgb, "bg": bg})
    return cover, kept


# ═════════════════════════════════════════════════════════════════
# 3. PDF ─ 주석 삭제 + 굳은 형광펜 덮기
# ═════════════════════════════════════════════════════════════════
_PAD = 0.7           # 덮는 네모를 이만큼(pt) 넓혀 가장자리에 색 실선이 남지 않게 한다
_MARGIN = 0.03       # 덮는 세기의 여유


def _cover_ops(items: list[dict]) -> str:
    """덮기 그림 명령. ① 채도 혼합으로 색을 뺀다 → ② 닷지(또는 곱하기) 혼합으로 바탕 밝기에 맞춘다
    → ③ 바탕에 색이 있으면 색 혼합으로 그 색조를 입힌다. 검은 글자는 세 단계 모두에서 검게 남는다.

    ②의 세기는 형광펜 밝기(Lh)와 바탕 밝기(Lb)로 정한다: 형광펜 회색 Lh 가 정확히 Lb 가 되게.
    이렇게 하면 글자 가장자리(글자와 형광펜이 섞인 점)도 "흰 바탕에 쓴 글자" 와 같은 값이 되어 글자가 가늘어지지 않는다.
    """
    def path(it):
        x0, y0, x1, y1 = it["rect"]
        return f"{x0 - _PAD:.2f} {y0 - _PAD:.2f} {x1 - x0 + 2 * _PAD:.2f} {y1 - y0 + 2 * _PAD:.2f} re f"

    ops = ["q /PMhlS gs 0.5 g"] + [path(it) for it in items] + ["Q"]
    dodge, mult, tint = [], [], []
    for it in items:
        lh, lb = _lum(it["rgb"]), _lum(it["bg"])
        if lb >= lh:
            # 0.03 의 여유: 색 관리 때문에 화면에 그려진 형광펜이 계산값보다 조금 어두울 수 있다 (옅은 회색 띠 방지)
            dodge.append(f"{min(0.97, max(0.0, 1 - lh / lb + _MARGIN)):.4f} g {path(it)}")
        else:
            mult.append(f"{lb / lh:.4f} g {path(it)}")
        r, g, b = it["bg"]
        if max(r, g, b) - min(r, g, b) > 0.03:
            tint.append(f"{r:.4f} {g:.4f} {b:.4f} rg {path(it)}")
    if dodge:
        ops += ["q /PMhlD gs"] + dodge + ["Q"]
    if mult:
        ops += ["q /PMhlM gs"] + mult + ["Q"]
    if tint:
        ops += ["q /PMhlC gs"] + tint + ["Q"]
    return "\n".join(ops)


def clean_pdf(writer, baked: dict[int, list[dict]], baked_kept: int, strip: bool, cover: bool) -> dict:
    """PdfWriter(원본을 복제한 사본)에서 형광펜 주석을 지우고 굳은 형광펜을 덮는다."""
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, StreamObject

    info = new_info(strip)
    n_baked = sum(len(v) for v in baked.values())
    for pno, page in enumerate(writer.pages, 1):
        arr = page.get("/Annots")
        if arr is None:
            continue
        arr = arr.get_object()
        for ref in list(arr):
            try:
                a = ref.get_object()
            except Exception:  # noqa: BLE001
                continue
            sub, title = a.get("/Subtype"), str(a.get("/T") or "")
            is_tool = title in TOOL_ANNOT_TITLES and sub in ("/Highlight", "/Text", "/Square")
            if sub != "/Highlight" and not is_tool:
                continue                  # 링크·댓글 등 다른 주석은 그대로
            info["tool_prev" if is_tool else "found"] += 1
            if strip:
                arr.remove(ref)
                if not is_tool:
                    info["removed"] += 1
    info["found"] += n_baked + baked_kept
    if not strip:
        return info
    info["kept"] += baked_kept
    if not cover:
        info["kept"] += n_baked
        return info

    def stream(data: str):
        st = StreamObject()
        st.set_data(data.encode("latin-1"))
        return writer._add_object(st)

    for pno, items in baked.items():
        if not items or pno > len(writer.pages):
            continue
        page = writer.pages[pno - 1]
        raw = page.raw_get("/Contents") if "/Contents" in page else None
        parts = []
        if raw is not None:
            obj = raw.get_object()
            parts = list(obj) if isinstance(obj, ArrayObject) else [raw]
        # 원래 내용은 q … Q 로 감싸 좌표계·색 상태가 덮기 그림에 새지 않게 한다
        page[NameObject("/Contents")] = ArrayObject([stream("q\n"), *parts, stream("\nQ\n"), stream(_cover_ops(items))])
        res = page.get("/Resources")
        res = res.get_object() if res is not None else DictionaryObject()
        ext = res.get("/ExtGState")
        ext = ext.get_object() if ext is not None else DictionaryObject()
        for name, mode in (("/PMhlS", "/Saturation"), ("/PMhlD", "/ColorDodge"), ("/PMhlM", "/Multiply"), ("/PMhlC", "/Color")):
            ext[NameObject(name)] = DictionaryObject({NameObject("/Type"): NameObject("/ExtGState"),
                                                      NameObject("/BM"): NameObject(mode)})
        res[NameObject("/ExtGState")] = ext
        page[NameObject("/Resources")] = res
        info["covered"] += len(items)
    return info
